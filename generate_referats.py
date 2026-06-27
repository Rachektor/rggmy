#!/usr/bin/env python3
"""Генератор рефератов и заданий в формате Word."""

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
import os

STUDENT = "Абулашвили Рамзес Тристанович"
SUPERVISOR = "доцент, к.и.н. Арапов С.В."
YEAR = "2026"
OUTPUT_DIR = "/workspace/referats"


def set_run_font(run, size=14, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold


def add_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, size=14, space_after=0):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.5
    pf.space_after = Pt(space_after)
    pf.first_line_indent = Cm(1.25) if align == WD_ALIGN_PARAGRAPH.JUSTIFY else Cm(0)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_heading(doc, text):
    return add_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, space_after=6)


def add_section_title(doc, text):
    p = add_paragraph(doc, text, bold=True, space_after=6)
    p.paragraph_format.first_line_indent = Cm(0)
    return p


def add_title_page(doc, discipline, topic):
    for text, align, size, bold in [
        ("Министерство науки и высшего образования РФ", WD_ALIGN_PARAGRAPH.CENTER, 14, False),
        ("Российский государственный гидрометеорологический университет", WD_ALIGN_PARAGRAPH.CENTER, 14, False),
        ("", WD_ALIGN_PARAGRAPH.CENTER, 14, False),
        ("Кафедра национальной безопасности и международного права", WD_ALIGN_PARAGRAPH.CENTER, 14, False),
        ("", WD_ALIGN_PARAGRAPH.CENTER, 14, False),
        ("", WD_ALIGN_PARAGRAPH.CENTER, 14, False),
        (f'Дисциплина: «{discipline}»', WD_ALIGN_PARAGRAPH.CENTER, 14, False),
        ("", WD_ALIGN_PARAGRAPH.CENTER, 14, False),
        ("Реферат к зачету", WD_ALIGN_PARAGRAPH.CENTER, 14, False),
        (f"Тема реферата: {topic}", WD_ALIGN_PARAGRAPH.CENTER, 14, False),
        ("", WD_ALIGN_PARAGRAPH.CENTER, 14, False),
        ("", WD_ALIGN_PARAGRAPH.CENTER, 14, False),
        ("", WD_ALIGN_PARAGRAPH.CENTER, 14, False),
        ("Подготовили:", WD_ALIGN_PARAGRAPH.LEFT, 14, False),
        ("студенты гр. ……..", WD_ALIGN_PARAGRAPH.LEFT, 14, False),
        (STUDENT, WD_ALIGN_PARAGRAPH.LEFT, 14, False),
        (f"Проверил:  {SUPERVISOR}", WD_ALIGN_PARAGRAPH.LEFT, 14, False),
        ("", WD_ALIGN_PARAGRAPH.CENTER, 14, False),
        ("", WD_ALIGN_PARAGRAPH.CENTER, 14, False),
        ("", WD_ALIGN_PARAGRAPH.CENTER, 14, False),
        ("Санкт-Петербург", WD_ALIGN_PARAGRAPH.CENTER, 14, False),
        (YEAR, WD_ALIGN_PARAGRAPH.CENTER, 14, False),
    ]:
        p = doc.add_paragraph()
        p.alignment = align
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold)
    doc.add_page_break()


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    return doc


def add_content(doc, sections):
    add_heading(doc, "Содержание")
    for title, _ in sections:
        add_paragraph(doc, title, align=WD_ALIGN_PARAGRAPH.LEFT)
        p = doc.paragraphs[-1]
        p.paragraph_format.first_line_indent = Cm(0)
    doc.add_page_break()

    for title, paragraphs in sections:
        add_section_title(doc, title)
        for para in paragraphs:
            add_paragraph(doc, para)


def save_document(discipline, topic, filename, sections):
    doc = setup_document()
    add_title_page(doc, discipline, topic)
    add_content(doc, sections)
    path = os.path.join(OUTPUT_DIR, filename)
    doc.save(path)
    return path


def expand_paragraphs(base_paragraphs, extra_blocks):
    result = list(base_paragraphs)
    for block in extra_blocks:
        result.extend(block)
    return result


# Общие блоки для достижения объёма
ECON_INTRO = [
    "Экономическая теория изучает закономерности функционирования рыночной системы и механизмы принятия решений экономическими субъектами. Понимание базовых категорий микро- и макроэкономики необходимо для анализа реальных рыночных процессов, оценки последствий государственного вмешательства и прогнозирования поведения потребителей и производителей.",
    "В современных условиях российская экономика сочетает элементы рыночной конкуренции с активной ролью государства. Это предопределяет повышенное значение анализа спроса и предложения, эластичности, рыночного равновесия и бюджетной политики. Теоретические положения, рассмотренные в настоящей работе, позволяют системно подойти к решению практических задач и сформировать целостное представление об исследуемой проблеме.",
    "Цель реферата — раскрыть теоретическое содержание темы, показать практическое значение рассматриваемых категорий и продемонстрировать умение применять экономический аппарат при решении конкретных задач. В работе использованы положения отечественных и зарубежных учебников по экономической теории, а также нормативные акты, регулирующие экономические отношения в Российской Федерации.",
]

LAW_INTRO = [
    "Право представляет собой систему общеобязательных норм, установленных или санкционированных государством и обеспеченных его принудительной силой. Юридическая наука и правоприменительная практика постоянно обращаются к анализу институтов, источников, субъектов и способов защиты прав, поскольку именно через эти категории реализуется правовое регулирование общественных отношений.",
    "В Российской Федерации правовая система формируется на основе Конституции, федеральных законов, подзаконных актов и иных источников права. Каждая отрасль права обладает собственным предметом, методом и системой институтов, однако все они объединены общими принципами справедливости, равенства перед законом и неотъемлемости прав и свобод человека.",
    "Цель настоящего реферата — всесторонне рассмотреть заданную тему, выявить её место в системе права, охарактеризовать основные понятия и показать значение изучаемых институтов для обеспечения законности и защиты прав граждан. При подготовке работы использованы Конституция РФ, кодексы, федеральные законы и учебная литература по соответствующим отраслям права.",
]

CONCLUSION_ECON = [
    "Таким образом, рассмотренная тема занимает важное место в системе экономической теории и имеет непосредственное практическое значение. Полученные теоретические выводы подтверждаются решением прикладных задач и позволяют более точно оценивать рыночные процессы.",
    "Знание основ микро- и макроэкономического анализа необходимо не только специалистам-экономистам, но и всем участникам рыночных отношений. Оно способствует принятию обоснованных решений, пониманию последствий изменения цен, доходов, налогов и государственных расходов.",
    "Дальнейшее изучение темы предполагает углублённый анализ эмпирических данных, сопоставление теоретических моделей с реальной экономической политикой и учёт институциональных особенностей российской экономики.",
]

CONCLUSION_LAW = [
    "Подводя итог, можно отметить, что рассмотренные правовые институты играют существенную роль в обеспечении стабильности правопорядка и защиты прав субъектов правоотношений. Их содержание непосредственно связано с конституционными принципами и требованиями законности.",
    "Эффективность правового регулирования зависит не только от качества нормативных актов, но и от уровня правовой культуры общества, квалификации правоприменителей и доступности механизмов судебной защиты. Поэтому изучение темы имеет как теоретическое, так и практическое значение.",
    "Совершенствование законодательства в данной сфере должно учитывать динамику общественных отношений, международные стандарты и потребности правоохранительной и судебной практики.",
]

REFERENCES_ECON = [
    "1. Налоговая политика и налогообложение в Российской Федерации: учебник / под ред. Н.И. Безруковой. — М.: Юрайт, 2023.",
    "2. Микроэкономика: учебник / под ред. Г.О. Громова. — М.: Юрайт, 2024.",
    "3. Макроэкономика: учебник / под ред. В.Е. Макарова. — М.: ИНФРА-М, 2023.",
    "4. Экономическая теория: учебник / под ред. А.С. Булатова. — М.: Юрайт, 2024.",
    "5. Рынок и рыночная экономика: учебное пособие / под ред. О.В. Григорьевой. — СПб.: СЗИУ РАНХиГС, 2022.",
]

REFERENCES_LAW = [
    "1. Конституция Российской Федерации (принята всенародным голосованием 12.12.1993).",
    "2. Теория государства и права: учебник / под ред. В.К. Бабаева. — М.: Юрайт, 2024.",
    "3. Гражданский кодекс Российской Федерации (часть первая) от 30.11.1994 № 51-ФЗ.",
    "4. Уголовный кодекс Российской Федерации от 13.06.1996 № 63-ФЗ.",
    "5. Кодекс Российской Федерации об административных правонарушениях от 30.12.2001 № 195-ФЗ.",
    "6. Административное право: учебник / под ред. Л.Л. Попова, М.С. Студеникиной. — М.: Проспект, 2024.",
]

if __name__ == "__main__":
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    print("Base module loaded. Run generate_all.py")
