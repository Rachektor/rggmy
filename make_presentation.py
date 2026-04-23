"""Генерация презентации "Присутственный этикет:
поведение в общественных местах, ресторанах, музеях и театрах" (15 слайдов).
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN


# ---------- Цветовая палитра ----------
NAVY = RGBColor(0x10, 0x2A, 0x54)     # тёмно-синий
GOLD = RGBColor(0xC8, 0xA2, 0x5B)     # золотой
CREAM = RGBColor(0xF7, 0xF2, 0xE7)    # кремовый фон
DARK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x55, 0x5B, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def add_background(slide, color=CREAM):
    """Заливка слайда фоном."""
    left = top = 0
    width = prs.slide_width
    height = prs.slide_height
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.line.fill.background()
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    # отправить назад
    spTree = shape._element.getparent()
    spTree.remove(shape._element)
    spTree.insert(2, shape._element)
    return shape


def add_accent_bar(slide, color=GOLD, height_in=0.12):
    """Горизонтальная золотая полоска-акцент сверху."""
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(height_in)
    )
    bar.line.fill.background()
    bar.fill.solid()
    bar.fill.fore_color.rgb = color


def add_side_bar(slide, color=NAVY, width_in=0.35):
    """Вертикальная тёмная полоса слева."""
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(width_in), prs.slide_height
    )
    bar.line.fill.background()
    bar.fill.solid()
    bar.fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, *,
                font_size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT,
                font_name="Calibri"):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)

    lines = text.split("\n") if isinstance(text, str) else list(text)
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = font_name
    return tb


def add_bullets(slide, left, top, width, height, items, *,
                font_size=18, color=DARK, bullet="•", line_spacing=1.25):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = f"{bullet}  {item}"
        run.font.size = Pt(font_size)
        run.font.color.rgb = color
        run.font.name = "Calibri"
    return tb


def add_slide_number(slide, num, total=15):
    add_textbox(
        slide,
        Inches(12.4), Inches(7.0), Inches(1.0), Inches(0.35),
        f"{num} / {total}",
        font_size=11, color=GREY, align=PP_ALIGN.RIGHT,
    )


def add_footer_brand(slide):
    add_textbox(
        slide,
        Inches(0.6), Inches(7.0), Inches(8.0), Inches(0.35),
        "Присутственный этикет · Правила хорошего тона",
        font_size=11, color=GREY,
    )


def add_title(slide, title_text, subtitle_text=None):
    add_accent_bar(slide)
    add_textbox(
        slide,
        Inches(0.7), Inches(0.45), Inches(12.0), Inches(0.9),
        title_text,
        font_size=34, bold=True, color=NAVY,
    )
    if subtitle_text:
        add_textbox(
            slide,
            Inches(0.7), Inches(1.25), Inches(12.0), Inches(0.55),
            subtitle_text,
            font_size=18, color=GOLD, bold=True,
        )
    # Разделитель
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.7), Inches(1.85), Inches(2.2), Inches(0.04),
    )
    line.line.fill.background()
    line.fill.solid()
    line.fill.fore_color.rgb = GOLD


def add_card(slide, left, top, width, height, title, items,
             title_color=NAVY, accent=GOLD):
    """Карточка с заголовком и списком пунктов."""
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    card.adjustments[0] = 0.06
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = accent
    card.line.width = Pt(1.25)

    # Заголовок карточки
    add_textbox(
        slide,
        left + Inches(0.25), top + Inches(0.15),
        width - Inches(0.5), Inches(0.5),
        title,
        font_size=18, bold=True, color=title_color,
    )
    # Подчёркивание
    ul = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        left + Inches(0.25), top + Inches(0.62),
        Inches(0.6), Inches(0.035),
    )
    ul.line.fill.background()
    ul.fill.solid()
    ul.fill.fore_color.rgb = accent

    add_bullets(
        slide,
        left + Inches(0.25), top + Inches(0.78),
        width - Inches(0.5), height - Inches(0.9),
        items,
        font_size=13, color=DARK, line_spacing=1.18,
    )


# ---------- Создание презентации ----------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

blank_layout = prs.slide_layouts[6]


# ====================================================================
# СЛАЙД 1 — Титульный
# ====================================================================
s = prs.slides.add_slide(blank_layout)
add_background(s, CREAM)

# Левая декоративная полоса
left_panel = s.shapes.add_shape(
    MSO_SHAPE.RECTANGLE, 0, 0, Inches(4.4), prs.slide_height
)
left_panel.line.fill.background()
left_panel.fill.solid()
left_panel.fill.fore_color.rgb = NAVY

# Золотая вертикальная полоска
gold_bar = s.shapes.add_shape(
    MSO_SHAPE.RECTANGLE, Inches(4.4), 0, Inches(0.08), prs.slide_height
)
gold_bar.line.fill.background()
gold_bar.fill.solid()
gold_bar.fill.fore_color.rgb = GOLD

# Декоративный текст на тёмной панели
add_textbox(
    s, Inches(0.6), Inches(0.6), Inches(3.5), Inches(0.5),
    "ETIQUETTE · 2026",
    font_size=14, bold=True, color=GOLD,
)
add_textbox(
    s, Inches(0.6), Inches(6.5), Inches(3.5), Inches(0.5),
    "Презентация",
    font_size=14, color=WHITE,
)

# Главный заголовок
add_textbox(
    s, Inches(5.0), Inches(1.6), Inches(7.8), Inches(1.6),
    "Присутственный\nэтикет",
    font_size=54, bold=True, color=NAVY,
)
# Разделитель
sep = s.shapes.add_shape(
    MSO_SHAPE.RECTANGLE, Inches(5.0), Inches(3.6), Inches(2.0), Inches(0.05),
)
sep.line.fill.background()
sep.fill.solid()
sep.fill.fore_color.rgb = GOLD

add_textbox(
    s, Inches(5.0), Inches(3.85), Inches(7.8), Inches(1.2),
    "Поведение в общественных местах,\nресторанах, музеях и театрах",
    font_size=22, color=DARK,
)

add_textbox(
    s, Inches(5.0), Inches(6.4), Inches(7.8), Inches(0.5),
    "15 слайдов · Правила хорошего тона",
    font_size=14, color=GREY, bold=True,
)


# ====================================================================
# СЛАЙД 2 — Что такое этикет
# ====================================================================
s = prs.slides.add_slide(blank_layout)
add_background(s)
add_title(s, "Что такое этикет?", "Введение в тему")

add_textbox(
    s, Inches(0.7), Inches(2.1), Inches(12.0), Inches(1.2),
    "Этикет — это совокупность общепринятых правил поведения, "
    "которые помогают людям взаимодействовать уважительно, "
    "доброжелательно и комфортно для окружающих.",
    font_size=18, color=DARK,
)

add_card(
    s, Inches(0.7), Inches(3.6), Inches(3.9), Inches(3.1),
    "Зачем нужен этикет",
    [
        "Создаёт уважительную атмосферу",
        "Помогает избегать конфликтов",
        "Формирует культуру общества",
        "Показывает уровень воспитания",
    ],
)
add_card(
    s, Inches(4.8), Inches(3.6), Inches(3.9), Inches(3.1),
    "Основные принципы",
    [
        "Уважение к другим людям",
        "Тактичность и сдержанность",
        "Опрятность и аккуратность",
        "Пунктуальность",
    ],
)
add_card(
    s, Inches(8.9), Inches(3.6), Inches(3.9), Inches(3.1),
    "Где он применим",
    [
        "Транспорт и улица",
        "Рестораны и кафе",
        "Музеи и выставки",
        "Театры и концерты",
    ],
)

add_footer_brand(s)
add_slide_number(s, 2)


# ====================================================================
# СЛАЙД 3 — История этикета
# ====================================================================
s = prs.slides.add_slide(blank_layout)
add_background(s)
add_title(s, "Немного истории", "От древности до наших дней")

timeline_items = [
    ("Древний мир",
     "Правила поведения при дворах фараонов и царей Востока."),
    ("Средневековье",
     "Рыцарский кодекс чести, придворный этикет Европы."),
    ("XVII–XVIII века",
     "При Людовике XIV появляется слово «étiquette» — карточки с правилами."),
    ("Россия, XVIII век",
     "Пётр I вводит «Юности честное зерцало» — первый русский учебник этикета."),
    ("Современность",
     "Этикет стал проще, но принцип уважения к людям остался неизменным."),
]

top = Inches(2.2)
for title, desc in timeline_items:
    # Маркер
    dot = s.shapes.add_shape(
        MSO_SHAPE.OVAL,
        Inches(0.9), top + Inches(0.12),
        Inches(0.22), Inches(0.22),
    )
    dot.line.fill.background()
    dot.fill.solid()
    dot.fill.fore_color.rgb = GOLD

    add_textbox(
        s, Inches(1.4), top, Inches(3.2), Inches(0.5),
        title, font_size=16, bold=True, color=NAVY,
    )
    add_textbox(
        s, Inches(4.6), top, Inches(8.3), Inches(0.6),
        desc, font_size=15, color=DARK,
    )
    top += Inches(0.85)

add_footer_brand(s)
add_slide_number(s, 3)


# ====================================================================
# СЛАЙД 4 — Общие правила поведения в общественных местах
# ====================================================================
s = prs.slides.add_slide(blank_layout)
add_background(s)
add_title(s, "Поведение в общественных местах",
          "Общие правила для всех ситуаций")

add_card(
    s, Inches(0.7), Inches(2.1), Inches(6.0), Inches(4.8),
    "Что делать",
    [
        "Соблюдать тишину и не мешать окружающим",
        "Уступать место пожилым, беременным, людям с детьми",
        "Здороваться и благодарить за помощь",
        "Следить за опрятным внешним видом",
        "Выбрасывать мусор только в урны",
        "Быть пунктуальным и предупредительным",
        "Извиняться, если задели кого-то случайно",
    ],
)
add_card(
    s, Inches(6.85), Inches(2.1), Inches(6.0), Inches(4.8),
    "Чего делать нельзя",
    [
        "Громко разговаривать и смеяться",
        "Говорить по телефону на повышенных тонах",
        "Курить в неположенных местах",
        "Толкаться, спешить, расталкивать людей",
        "Сорить, плевать, жевать жвачку на ходу",
        "Показывать пальцем, обсуждать прохожих",
        "Проявлять агрессию и грубость",
    ],
    accent=RGBColor(0xB0, 0x3A, 0x2E),
)

add_footer_brand(s)
add_slide_number(s, 4)


# ====================================================================
# СЛАЙД 5 — Общественный транспорт
# ====================================================================
s = prs.slides.add_slide(blank_layout)
add_background(s)
add_title(s, "В транспорте и на улице",
          "Как вести себя среди людей")

add_bullets(
    s, Inches(0.8), Inches(2.1), Inches(11.8), Inches(5.0),
    [
        "Пропускаем выходящих из транспорта, а затем заходим сами.",
        "Уступаем места тем, кому они нужнее: пожилым, беременным, инвалидам, пассажирам с детьми.",
        "Не едим в транспорте пищу с резким запахом и не занимаем сиденья сумками.",
        "Слушаем музыку только в наушниках и на умеренной громкости.",
        "В разговоре не обсуждаем личные темы громко — соседи слышат всё.",
        "На эскалаторе стоим справа, проход слева оставляем для торопящихся.",
        "На улице не останавливаемся внезапно в потоке людей — отходим в сторону.",
        "Дверь придерживаем для того, кто идёт следом, особенно для пожилых людей.",
    ],
    font_size=18, line_spacing=1.35,
)

add_footer_brand(s)
add_slide_number(s, 5)


# ====================================================================
# СЛАЙД 6 — Этикет в ресторане: до еды
# ====================================================================
s = prs.slides.add_slide(blank_layout)
add_background(s)
add_title(s, "Этикет в ресторане: прибытие",
          "С чего начинается визит")

add_bullets(
    s, Inches(0.8), Inches(2.1), Inches(11.8), Inches(5.0),
    [
        "Столик лучше бронировать заранее, а в случае отмены — предупредить заведение.",
        "Приходить желательно вовремя; опоздание более чем на 10–15 минут требует извинений.",
        "Верхнюю одежду и крупные сумки оставляют в гардеробе.",
        "Мужчина помогает спутнице снять пальто и отодвигает стул.",
        "За столом женщина садится первой, мужчина — после.",
        "Телефон кладут экраном вниз или убирают в сумку; звонок — только по необходимости.",
        "Меню изучают спокойно; если возникли вопросы — обращаются к официанту вежливо.",
        "Заказ делают по очереди: сначала дамы, затем мужчины.",
    ],
    font_size=17, line_spacing=1.3,
)

add_footer_brand(s)
add_slide_number(s, 6)


# ====================================================================
# СЛАЙД 7 — Этикет за столом
# ====================================================================
s = prs.slides.add_slide(blank_layout)
add_background(s)
add_title(s, "Этикет за столом", "Как правильно вести себя во время еды")

add_card(
    s, Inches(0.7), Inches(2.1), Inches(6.0), Inches(4.8),
    "Хорошие манеры",
    [
        "Салфетку кладут на колени, развернув пополам",
        "Приборы используют «от краёв к центру»",
        "Едят бесшумно, с закрытым ртом",
        "Говорят только после того, как проглотили пищу",
        "Благодарят официанта за обслуживание",
        "Локти не ставят на стол во время еды",
        "Нож держат в правой, вилку — в левой руке",
    ],
)
add_card(
    s, Inches(6.85), Inches(2.1), Inches(6.0), Inches(4.8),
    "Чего не делают",
    [
        "Не разговаривают с набитым ртом",
        "Не машут приборами и не указывают ими на людей",
        "Не дуют на горячее блюдо — ждут, пока остынет",
        "Не чавкают, не хлюпают, не стучат приборами",
        "Не делают замечаний спутникам громко",
        "Не кладут телефон на стол рядом с тарелкой",
        "Не тянутся через весь стол — просят передать",
    ],
    accent=RGBColor(0xB0, 0x3A, 0x2E),
)

add_footer_brand(s)
add_slide_number(s, 7)


# ====================================================================
# СЛАЙД 8 — Сервировка и приборы
# ====================================================================
s = prs.slides.add_slide(blank_layout)
add_background(s)
add_title(s, "Сервировка и столовые приборы",
          "Маленькая азбука застолья")

add_bullets(
    s, Inches(0.8), Inches(2.1), Inches(11.8), Inches(5.0),
    [
        "Приборы используют последовательно — от крайних к ближним к тарелке.",
        "Вилка слева, нож и ложка — справа. Десертные приборы расположены над тарелкой.",
        "Бокал для воды ставят ближе к тарелке, для вина — чуть дальше.",
        "Если делают паузу: приборы кладут крест-накрест или «на 8 часов» на тарелке.",
        "Когда закончили есть: приборы кладут параллельно, ручками вправо — «на 4 часа».",
        "Хлеб не режут ножом, а отламывают небольшими кусочками.",
        "Салфетку после трапезы кладут слева от тарелки, слегка смятой.",
    ],
    font_size=18, line_spacing=1.35,
)

add_footer_brand(s)
add_slide_number(s, 8)


# ====================================================================
# СЛАЙД 9 — Оплата счёта и чаевые
# ====================================================================
s = prs.slides.add_slide(blank_layout)
add_background(s)
add_title(s, "Счёт, чаевые и прощание",
          "Завершение визита в ресторан")

add_bullets(
    s, Inches(0.8), Inches(2.1), Inches(11.8), Inches(5.0),
    [
        "Счёт просят жестом или вежливой фразой: «Счёт, пожалуйста».",
        "Оплачивает обычно тот, кто пригласил; в компании друзей — делят поровну или каждый за себя.",
        "Чаевые — знак благодарности за обслуживание, обычно 10% от суммы счёта.",
        "Если обслуживание не понравилось, замечание делают спокойно и корректно.",
        "Перед уходом благодарят официанта и прощаются с администратором.",
        "Стул при уходе аккуратно задвигают к столу.",
        "Повышать голос, выяснять отношения или критиковать заведение громко — недопустимо.",
    ],
    font_size=18, line_spacing=1.35,
)

add_footer_brand(s)
add_slide_number(s, 9)


# ====================================================================
# СЛАЙД 10 — Этикет в музее
# ====================================================================
s = prs.slides.add_slide(blank_layout)
add_background(s)
add_title(s, "Этикет в музее", "Как вести себя среди экспонатов")

add_card(
    s, Inches(0.7), Inches(2.1), Inches(6.0), Inches(4.8),
    "Рекомендуется",
    [
        "Сдать верхнюю одежду и крупные сумки в гардероб",
        "Говорить тихо, вполголоса",
        "Соблюдать дистанцию между экспонатами",
        "Внимательно читать таблички и слушать экскурсовода",
        "Уступать место у экспонатов другим посетителям",
        "Задавать вопросы экскурсоводу в подходящий момент",
        "Отключить звук телефона",
    ],
)
add_card(
    s, Inches(6.85), Inches(2.1), Inches(6.0), Inches(4.8),
    "Недопустимо",
    [
        "Трогать экспонаты руками",
        "Перегибаться через ограждения",
        "Фотографировать со вспышкой (если запрещено)",
        "Есть, пить и жевать жвачку в залах",
        "Громко разговаривать и смеяться",
        "Бегать по залам и обгонять экскурсию",
        "Опираться на стены, витрины и постаменты",
    ],
    accent=RGBColor(0xB0, 0x3A, 0x2E),
)

add_footer_brand(s)
add_slide_number(s, 10)


# ====================================================================
# СЛАЙД 11 — Музей: на экскурсии
# ====================================================================
s = prs.slides.add_slide(blank_layout)
add_background(s)
add_title(s, "На экскурсии и выставке",
          "Уважение к искусству и к другим посетителям")

add_bullets(
    s, Inches(0.8), Inches(2.1), Inches(11.8), Inches(5.0),
    [
        "Держимся рядом с группой, не отстаём и не убегаем вперёд.",
        "Слушаем экскурсовода внимательно, не перебиваем — вопросы задаём в конце.",
        "Если пришли с детьми — заранее объясните им правила поведения.",
        "Не обсуждаем экспонаты громко и не критикуем работы авторов у всех на виду.",
        "Перед фотографией убеждаемся, что съёмка разрешена, и не загораживаем других.",
        "Уважаем труд смотрителей: их просьбы — не придирки, а правила безопасности.",
        "Дресс-код в музее свободный, но одежда должна быть чистой и уместной.",
    ],
    font_size=18, line_spacing=1.35,
)

add_footer_brand(s)
add_slide_number(s, 11)


# ====================================================================
# СЛАЙД 12 — Этикет в театре: до спектакля
# ====================================================================
s = prs.slides.add_slide(blank_layout)
add_background(s)
add_title(s, "Этикет в театре: до спектакля",
          "Подготовка к культурному вечеру")

add_bullets(
    s, Inches(0.8), Inches(2.1), Inches(11.8), Inches(5.0),
    [
        "В театр приходят нарядно: классический стиль, опрятная одежда и обувь.",
        "Приходить принято за 20–30 минут до начала спектакля.",
        "Верхнюю одежду обязательно сдают в гардероб, номерок убирают надёжно.",
        "Перед входом в зал посещают фойе, знакомятся с программкой.",
        "К своему месту проходят лицом к сидящим, извиняясь за беспокойство.",
        "Мужчина идёт по ряду первым, освобождая дорогу даме.",
        "Мобильный телефон выключают или переводят в беззвучный режим.",
        "Если опоздали — проходят только после первой сцены с разрешения билетёра.",
    ],
    font_size=17, line_spacing=1.3,
)

add_footer_brand(s)
add_slide_number(s, 12)


# ====================================================================
# СЛАЙД 13 — Этикет в театре: во время спектакля
# ====================================================================
s = prs.slides.add_slide(blank_layout)
add_background(s)
add_title(s, "Во время и после спектакля",
          "Уважение к актёрам и зрителям")

add_card(
    s, Inches(0.7), Inches(2.1), Inches(6.0), Inches(4.8),
    "Во время действия",
    [
        "Сидим тихо, не разговариваем и не комментируем",
        "Не шуршим пакетами, программками, фантиками",
        "Не едим и не пьём в зрительном зале",
        "Не фотографируем и не снимаем видео",
        "Не раскачиваемся и не наклоняемся вперёд",
        "Если стало плохо — тихо выходим в боковой проход",
    ],
)
add_card(
    s, Inches(6.85), Inches(2.1), Inches(6.0), Inches(4.8),
    "В антракте и после",
    [
        "В антракте не спешим — выходим спокойно",
        "В буфете соблюдаем очередь, говорим тихо",
        "Аплодируем искренне, но сдержанно",
        "Крики «браво!» — для настоящего восхищения",
        "Цветы артистам дарят в конце спектакля",
        "Уходить до финальных поклонов — неуважение",
    ],
)

add_footer_brand(s)
add_slide_number(s, 13)


# ====================================================================
# СЛАЙД 14 — Золотые правила
# ====================================================================
s = prs.slides.add_slide(blank_layout)
add_background(s)
add_title(s, "Золотые правила этикета",
          "Коротко — о самом главном")

rules = [
    ("01", "Уважай других", "Относись к людям так, как хочешь, чтобы относились к тебе."),
    ("02", "Будь тактичен", "Не ставь окружающих в неловкое положение."),
    ("03", "Соблюдай тишину", "Громкость — главный враг общественных мест."),
    ("04", "Следи за собой", "Опрятность и пунктуальность — основа уважения."),
    ("05", "Благодари", "«Спасибо» и «пожалуйста» открывают любые двери."),
    ("06", "Знай место", "В ресторане, музее и театре — свои правила. Уточни их заранее."),
]

left_start = Inches(0.7)
top_start = Inches(2.1)
card_w = Inches(4.0)
card_h = Inches(2.25)
gap_x = Inches(0.15)
gap_y = Inches(0.2)

for i, (num, title, desc) in enumerate(rules):
    row, col = divmod(i, 3)
    left = left_start + col * (card_w + gap_x)
    top = top_start + row * (card_h + gap_y)
    card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, card_w, card_h)
    card.adjustments[0] = 0.08
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = GOLD
    card.line.width = Pt(1.25)

    add_textbox(
        s, left + Inches(0.2), top + Inches(0.1),
        Inches(1.0), Inches(0.6),
        num, font_size=30, bold=True, color=GOLD,
    )
    add_textbox(
        s, left + Inches(1.2), top + Inches(0.18),
        card_w - Inches(1.3), Inches(0.55),
        title, font_size=17, bold=True, color=NAVY,
    )
    add_textbox(
        s, left + Inches(0.25), top + Inches(0.95),
        card_w - Inches(0.5), card_h - Inches(1.1),
        desc, font_size=13, color=DARK,
    )

add_footer_brand(s)
add_slide_number(s, 14)


# ====================================================================
# СЛАЙД 15 — Заключение / Спасибо за внимание
# ====================================================================
s = prs.slides.add_slide(blank_layout)
add_background(s, NAVY)

# Золотая рамка внутри
frame = s.shapes.add_shape(
    MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(0.5),
    prs.slide_width - Inches(1.0), prs.slide_height - Inches(1.0),
)
frame.fill.background()
frame.line.color.rgb = GOLD
frame.line.width = Pt(1.5)

add_textbox(
    s, Inches(0.7), Inches(1.5), Inches(12.0), Inches(0.8),
    "ЗАКЛЮЧЕНИЕ",
    font_size=16, bold=True, color=GOLD, align=PP_ALIGN.CENTER,
)

add_textbox(
    s, Inches(0.7), Inches(2.3), Inches(12.0), Inches(1.4),
    "Спасибо за внимание!",
    font_size=54, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
)

# Разделитель
sep = s.shapes.add_shape(
    MSO_SHAPE.RECTANGLE,
    Inches(6.17), Inches(3.9), Inches(1.0), Inches(0.05),
)
sep.line.fill.background()
sep.fill.solid()
sep.fill.fore_color.rgb = GOLD

add_textbox(
    s, Inches(1.0), Inches(4.2), Inches(11.3), Inches(1.6),
    "Этикет — это не набор запретов, а способ сделать мир вокруг "
    "немного добрее и красивее.\nСоблюдая простые правила, мы "
    "проявляем уважение к себе и к людям.",
    font_size=20, color=CREAM, align=PP_ALIGN.CENTER,
)

add_textbox(
    s, Inches(1.0), Inches(6.2), Inches(11.3), Inches(0.6),
    "Ведите себя достойно — и мир ответит тем же.",
    font_size=16, bold=True, color=GOLD, align=PP_ALIGN.CENTER,
)

# Итог: сохранение
out = "Присутственный_этикет.pptx"
prs.save(out)
print(f"Готово: {out}")
