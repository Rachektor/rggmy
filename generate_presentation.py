"""Generate a 15-slide PowerPoint presentation on travel etiquette (in Russian)."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN


PRIMARY = RGBColor(0x0F, 0x3D, 0x6E)       # тёмно-синий
ACCENT = RGBColor(0xE0, 0x8E, 0x2A)        # тёплый оранжевый
LIGHT = RGBColor(0xF5, 0xF1, 0xE8)         # слоновая кость
TEXT_DARK = RGBColor(0x22, 0x2B, 0x3A)
TEXT_MUTED = RGBColor(0x55, 0x5D, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


SLIDES = [
    {
        "type": "title",
        "title": "Этикет путешествий",
        "subtitle": "Как вести себя культурно, уважительно и с удовольствием в любой точке мира",
    },
    {
        "type": "content",
        "title": "О чём эта презентация",
        "bullets": [
            "Что такое этикет путешествий и зачем он нужен",
            "Правила поведения до, во время и после поездки",
            "Уважение к местным традициям и культуре",
            "Этикет в транспорте, отелях и общественных местах",
            "Экологичные и ответственные путешествия",
        ],
    },
    {
        "type": "content",
        "title": "Что такое этикет путешественника",
        "bullets": [
            "Набор негласных правил поведения вдали от дома",
            "Уважение к стране, людям и их образу жизни",
            "Забота о попутчиках, природе и культурном наследии",
            "Визитная карточка вашей страны за рубежом",
        ],
    },
    {
        "type": "content",
        "title": "Подготовка к поездке",
        "bullets": [
            "Изучите культуру, религию и обычаи страны",
            "Выучите базовые фразы: «здравствуйте», «спасибо», «извините»",
            "Ознакомьтесь с законами и запретами",
            "Уточните правила чаевых и дресс-кода",
            "Проверьте документы, страховку и прививки",
        ],
    },
    {
        "type": "content",
        "title": "Этикет в аэропорту и самолёте",
        "bullets": [
            "Приезжайте заранее и не занимайте чужое время в очередях",
            "Не разговаривайте громко и уважайте личное пространство",
            "Откидывайте кресло аккуратно и только при необходимости",
            "Делите подлокотник и не занимайте чужой багажный отсек",
            "Будьте вежливы с экипажем — это их рабочее место",
        ],
    },
    {
        "type": "content",
        "title": "Поведение в поезде и автобусе",
        "bullets": [
            "Не ешьте резко пахнущую еду в закрытом пространстве",
            "Используйте наушники при прослушивании музыки и видео",
            "Уступайте места пожилым, детям и людям с ограничениями",
            "Держите багаж компактно, не загромождайте проходы",
            "Разговаривайте вполголоса, особенно в ночное время",
        ],
    },
    {
        "type": "content",
        "title": "Этикет в отеле",
        "bullets": [
            "Здоровайтесь с персоналом и благодарите за помощь",
            "Соблюдайте тишину после 22:00",
            "Не выносите из номера посуду и предметы интерьера",
            "Оставляйте чаевые горничным и носильщикам",
            "При выезде оставляйте номер в приличном виде",
        ],
    },
    {
        "type": "content",
        "title": "За столом: ресторанный этикет",
        "bullets": [
            "Изучите местные традиции еды: палочки, руки, приборы",
            "Не критикуйте национальную кухню и способ подачи",
            "Уточните, принято ли делить счёт и оставлять чаевые",
            "Не фотографируйте еду и персонал без разрешения",
            "Будьте терпеливы: в разных странах — разный темп обслуживания",
        ],
    },
    {
        "type": "content",
        "title": "Уважение к местным обычаям",
        "bullets": [
            "Одевайтесь скромно в храмах, мечетях и монастырях",
            "Снимайте обувь там, где это принято",
            "Не трогайте сакральные предметы и статуи",
            "Уважайте время молитвы и религиозные праздники",
            "Избегайте громких оценок и сравнений с домом",
        ],
    },
    {
        "type": "content",
        "title": "Фото- и видеоэтикет",
        "bullets": [
            "Спрашивайте разрешение, прежде чем снимать людей",
            "Не фотографируйте военные и государственные объекты",
            "Уважайте запреты на съёмку в музеях и храмах",
            "Не мешайте другим туристам ради идеального кадра",
            "Публикуйте деликатно — уважая частную жизнь",
        ],
    },
    {
        "type": "content",
        "title": "Общение с местными жителями",
        "bullets": [
            "Улыбайтесь и здоровайтесь на местном языке",
            "Говорите спокойно, не повышайте голос при недопонимании",
            "Не навязывайте свои взгляды на политику и религию",
            "Будьте открытыми, но сохраняйте разумную осторожность",
            "Благодарите за помощь — искренне и вовремя",
        ],
    },
    {
        "type": "content",
        "title": "Экологичные путешествия",
        "bullets": [
            "Берите многоразовую бутылку и сумку",
            "Сортируйте мусор и не оставляйте его в природе",
            "Не кормите диких животных и не нарушайте их покой",
            "Выбирайте местных гидов и небольшие эко-отели",
            "Экономьте воду и электроэнергию в номере",
        ],
    },
    {
        "type": "content",
        "title": "Безопасность и ответственность",
        "bullets": [
            "Храните копии документов отдельно от оригиналов",
            "Не демонстрируйте крупные суммы и дорогую технику",
            "Соблюдайте местные законы даже в мелочах",
            "Имейте контакты посольства и страховой компании",
            "Уважайте правила дорожного движения страны",
        ],
    },
    {
        "type": "content",
        "title": "Типичные ошибки туристов",
        "bullets": [
            "Громкое поведение в общественных местах",
            "Сравнение «у нас лучше» и критика местных порядков",
            "Игнорирование дресс-кода в храмах",
            "Торг там, где он неуместен, и наоборот",
            "Неуважение к очередям и личному пространству",
        ],
    },
    {
        "type": "final",
        "title": "Путешествуйте с уважением",
        "subtitle": "Хороший турист оставляет после себя только добрые воспоминания",
        "bullets": [
            "Уважайте страну — и страна ответит тем же",
            "Будьте гостем, а не ревизором",
            "Учитесь у мира — и делитесь своей культурой бережно",
        ],
    },
]


def set_bg(slide, color):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.line.fill.background()
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.shadow.inherit = False
    slide.shapes._spTree.remove(bg._element)
    slide.shapes._spTree.insert(2, bg._element)
    return bg


def add_rect(slide, left, top, width, height, color, line=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if not line:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_text(slide, left, top, width, height, text, *, size=18, bold=False,
             color=TEXT_DARK, align=PP_ALIGN.LEFT, font="Calibri"):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tb


def add_bullets(slide, left, top, width, height, bullets, *, size=20, color=TEXT_DARK):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    for i, text in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = f"•  {text}"
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return tb


def build_title_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, PRIMARY)
    add_rect(slide, Inches(0), Inches(3.0), Inches(13.333), Inches(1.5), ACCENT)
    add_text(slide, Inches(0.8), Inches(1.2), Inches(11.7), Inches(1.2),
             "ЭТИКЕТ", size=28, bold=True, color=LIGHT)
    add_text(slide, Inches(0.8), Inches(1.75), Inches(11.7), Inches(2.0),
             data["title"].upper(), size=60, bold=True, color=WHITE)
    add_text(slide, Inches(0.8), Inches(4.7), Inches(11.7), Inches(1.0),
             data["subtitle"], size=22, color=LIGHT)
    add_text(slide, Inches(0.8), Inches(6.6), Inches(11.7), Inches(0.5),
             "Презентация • 15 слайдов", size=14, color=LIGHT)


def build_content_slide(prs, data, index, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, LIGHT)
    add_rect(slide, Inches(0), Inches(0), Inches(0.35), Inches(7.5), PRIMARY)
    add_rect(slide, Inches(0.35), Inches(0), Inches(12.98), Inches(1.35), PRIMARY)
    add_rect(slide, Inches(0.35), Inches(1.35), Inches(12.98), Inches(0.12), ACCENT)

    add_text(slide, Inches(0.8), Inches(0.35), Inches(11.7), Inches(0.6),
             f"Слайд {index} / {total}", size=14, color=LIGHT)
    add_text(slide, Inches(0.8), Inches(0.65), Inches(11.7), Inches(0.8),
             data["title"], size=30, bold=True, color=WHITE)

    add_bullets(slide, Inches(0.9), Inches(1.9), Inches(11.5), Inches(5.2),
                data["bullets"], size=22, color=TEXT_DARK)

    add_text(slide, Inches(0.8), Inches(7.05), Inches(11.7), Inches(0.35),
             "Этикет путешествий", size=11, color=TEXT_MUTED)


def build_final_slide(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, PRIMARY)
    add_rect(slide, Inches(0), Inches(6.0), Inches(13.333), Inches(0.2), ACCENT)
    add_text(slide, Inches(0.8), Inches(1.0), Inches(11.7), Inches(1.2),
             data["title"].upper(), size=48, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, Inches(0.8), Inches(2.2), Inches(11.7), Inches(1.0),
             data["subtitle"], size=22, color=LIGHT, align=PP_ALIGN.CENTER)

    box = slide.shapes.add_textbox(Inches(2.2), Inches(3.4), Inches(9), Inches(3))
    tf = box.text_frame
    tf.word_wrap = True
    for i, b in enumerate(data["bullets"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(14)
        run = p.add_run()
        run.text = f"✓  {b}"
        run.font.size = Pt(22)
        run.font.color.rgb = WHITE
        run.font.name = "Calibri"

    add_text(slide, Inches(0.8), Inches(6.8), Inches(11.7), Inches(0.4),
             "Спасибо за внимание!", size=18, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    total = len(SLIDES)
    for i, data in enumerate(SLIDES, start=1):
        if data["type"] == "title":
            build_title_slide(prs, data)
        elif data["type"] == "final":
            build_final_slide(prs, data)
        else:
            build_content_slide(prs, data, i, total)

    out = "Этикет_путешествий.pptx"
    prs.save(out)
    print(f"Saved: {out} ({total} slides)")


if __name__ == "__main__":
    main()
